from sqlalchemy import select, insert, create_engine, text, MetaData
from datetime import datetime
import openpyxl


host = "localhost"
port = 5432
username = "postgres"
password = "Bghujknmol123"
db = "demo"

engine = create_engine(f"postgresql+psycopg2://{username}:{password}@{host}:{port}/{db}")

metadata = MetaData()
metadata.reflect(bind=engine)

User = metadata.tables["user"]
Delivery = metadata.tables["delivery"]
Order = metadata.tables["order"]
OrderItem = metadata.tables["order_item"]
Product = metadata.tables["product"]


def normalize_article(value: str):
    articles = []
    quantities = []
    raw = [x.strip() for x in value.split(",")]

    for i in range(0, len(raw), 2):
        articles.append(raw[i])
        quantities.append(raw[i+1])
    return articles, quantities

if __name__ == "__main__":
    with engine.begin() as conn:
        conn.execute(text('TRUNCATE "user", "delivery", "order_item", "order", "product" CASCADE;'))


    user_table = openpyxl.load_workbook("import/user_import.xlsx").active

    user_rows = []
    for row in user_table.iter_rows(values_only=True, min_row=2):
        if row[0] is None:
            continue
        user_rows.append({
            "role": row[0],
            "name": row[1],
            "login": row[2],
            "password": row[3]
        })

    with engine.begin() as conn:
        conn.execute(insert(User), user_rows)
    
    with engine.begin() as conn:
        user_map = dict(conn.execute(select(User.c.name, User.c.id).where(User.c.role != "Авторизованный пользователь")).fetchall())

    delivery_table = openpyxl.load_workbook("import/Пункты выдачи_import.xlsx").active

    delivery_rows = []
    for i, row in enumerate(delivery_table.iter_rows(values_only=True, min_row=1)):
        if row[0] is None:
            continue
        delivery_rows.append({
            "id": i+1,
            "address": row[0]
        })

    with engine.begin() as conn:
        conn.execute(insert(Delivery), delivery_rows)
    
    product_table = openpyxl.load_workbook("import/Tovar.xlsx").active

    product_rows = []
    for row in product_table.iter_rows(values_only=True, min_row=2):
        if row[0] is None:
            continue
        product_rows.append({
            "article": row[0],
            "name": row[1],
            "measure_type": row[2],
            "price": row[3],
            "supplier": row[4],
            "producer": row[5],
            "category": row[6],
            "discount": row[7],
            "quantity": row[8],
            "description": row[9],
            "image_url": row[10]
        })

    with engine.begin() as conn:
        conn.execute(insert(Product), product_rows)


    order_table = openpyxl.load_workbook("import/Заказ_import.xlsx").active

    order_rows = []
    order_item_rows = []

    for row in order_table.iter_rows(values_only=True, min_row=2):
        if row[0] is None:
            continue
        
        order_id = row[0]
        article_raw = row[1]
        order_date = row[2]
        delivery_date = row[3]
        delivery_id = row[4]
        user_raw = row[5]
        challenge_code = row[6]
        status = row[7]

        if not isinstance(delivery_date, datetime) or not isinstance(order_date, datetime):
            continue

        order_rows.append({
            "id": order_id,
            "order_date": order_date,
            "delivery_date": delivery_date,
            "address_id": delivery_id,
            "user_id": user_map.get(user_raw),
            "challenge_code": challenge_code,
            "status": status
        })

        articles, quantities = normalize_article(article_raw)
        for a, q in zip(articles, quantities):
            order_item_rows.append({
                "order_id": order_id,
                "article": a,
                "quantity": q
            })

    with engine.begin() as conn:
        conn.execute(insert(Order), order_rows)
        conn.execute(insert(OrderItem), order_item_rows)
        
